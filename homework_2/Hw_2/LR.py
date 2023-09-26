import os
import zipfile
import PyPDF2
import openpyxl
from docx import Document
from PIL import Image
import pytesseract

# 定义一个函数，用于检查文件是否包含关键词
def file_contains_keyword(file_path, keyword):
    if file_path.endswith(".pdf"):
        try:
            pdf_file = open(file_path, 'rb')
            pdf_reader = PyPDF2.PdfFileReader(pdf_file)
            for page_num in range(pdf_reader.numPages):
                page = pdf_reader.getPage(page_num)
                page_text = page.extractText()
                if keyword in page_text:
                    return True
        except Exception as e:
            pass
        finally:
            pdf_file.close()

    elif file_path.endswith((".xlsx", ".xls")):
        try:
            excel_workbook = openpyxl.load_workbook(file_path, read_only=True)
            for sheet in excel_workbook:
                for row in sheet.iter_rows(values_only=True):
                    for cell_value in row:
                        if isinstance(cell_value, str) and keyword in cell_value:
                            return True
        except Exception as e:
            pass

    elif file_path.endswith(".docx"):
        try:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                if keyword in paragraph.text:
                    return True
        except Exception as e:
            pass

    elif file_path.endswith((".jpg", ".jpeg", ".png", ".gif")):
        try:
            img = Image.open(file_path)
            img_text = pytesseract.image_to_string(img)
            if keyword in img_text:
                return True
        except Exception as e:
            pass

    else:
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()
                if keyword in content:
                    return True
        except Exception as e:
            pass

    return False

# 定义一个函数，用于递归搜索文件夹中的文件
def search_files_with_keyword(directory, keyword):
    found_files = []

    for root, _, files in os.walk(directory):
        for file in files:
            file_path = os.path.join(root, file)
            if file_contains_keyword(file_path, keyword):
                found_files.append(file_path)

    return found_files

if __name__ == "__main__":
    folder_path = "综测材料"  # 替换为你的文件夹路径
    search_keyword = "王志涛"  # 替换为你要搜索的关键词

    found_files = search_files_with_keyword(folder_path, search_keyword)

    if found_files:
        print(f"包含关键词 '{search_keyword}' 的文件：")
        for file_path in found_files:
            print(file_path)
    else:
        print(f"未找到包含关键词 '{search_keyword}' 的文件。")
